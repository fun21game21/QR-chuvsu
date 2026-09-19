from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from uuid import UUID

import pytest
from conftest import student_data
from openpyxl import load_workbook
from PIL import Image


def attend(browser, session, **changes):
    return browser.post(f"/api/attendance/{session['session_code']}", student_data(**changes))


def test_new_student_and_journal(browser, teacher, session):
    result = attend(browser(), session)
    assert result.status_code == 200
    assert result.json()["message"] == "Отметка подтверждена"
    rows = teacher.get(f"/api/sessions/{session['id']}/attendance").json()
    assert len(rows) == 1 and rows[0]["status"] == "accepted"
    assert "device_token" not in rows[0]
    assert "Иванов Иван Иванович" in teacher.get(session["teacher_url"]).text


def test_duplicate_normalized(browser, teacher, session):
    student = browser()
    assert attend(student, session).status_code == 200
    result = attend(student, session, name="  иванов   иван  иванович ", group_name="ивт-21")
    assert result.status_code == 409
    assert result.json()["message"] == "Вы уже отметились"
    rows = teacher.get(f"/api/sessions/{session['id']}/attendance").json()
    assert [r["status"] for r in rows].count("accepted") == 1
    assert len(rows) == 2


def test_expired_session(client, browser, session):
    client.portal.call(
        client.app.state.pool.execute,
        "UPDATE sessions SET created_at=now()-interval '2 hours', "
        "expires_at=now()-interval '1 hour' WHERE id=$1",
        UUID(session["id"]),
    )
    result = attend(browser(), session)
    assert result.status_code == 410
    assert result.json()["message"] == "Сессия недействительна"
    assert browser().get(session["student_url"]).status_code == 410


@pytest.mark.parametrize(
    "coords,reason",
    [
        ({"latitude": None, "longitude": None}, "no_location"),
        ({"latitude": None}, "no_location"),
        ({"latitude": 55.75, "longitude": 37.61}, "outside_zone"),
    ],
)
def test_geolocation_rejections(browser, teacher, session, coords, reason):
    result = attend(browser(), session, **coords)
    assert result.status_code == 422
    assert result.json()["reason"] == reason
    rows = teacher.get(f"/api/sessions/{session['id']}/attendance").json()
    assert rows[0]["status"] == "rejected"
    assert rows[0]["reason"] == result.json()["message"]


def test_no_location_does_not_consume_device(browser, session):
    student = browser()
    assert attend(student, session, latitude=None).status_code == 422
    assert attend(student, session).status_code == 200


def test_same_device_other_name(browser, session):
    student = browser()
    assert attend(student, session).status_code == 200
    result = attend(student, session, name="Петров Пётр Петрович", seed="different")
    assert result.status_code == 409
    assert result.json()["reason"] == "device_used"


def test_same_fingerprint_new_cookie(browser, session):
    assert attend(browser(), session).status_code == 200
    result = attend(browser(), session, name="Петров Пётр Петрович")
    assert result.status_code == 409
    assert result.json()["reason"] == "fingerprint_used"


def test_same_ip_different_devices_allowed(browser, session):
    assert attend(browser(), session).status_code == 200
    assert attend(browser(), session, name="Петров Пётр Петрович", seed="two").status_code == 200


def test_missing_fingerprint(browser, session):
    assert attend(browser(), session, fingerprint=None).json()["reason"] == "no_fingerprint"


def test_wrong_group(browser, session):
    assert attend(browser(), session, group_name="ДРУГАЯ").json()["reason"] == "wrong_group"


def test_manual_and_closed_session(browser, teacher, session):
    path = f"/api/sessions/{session['id']}"
    manual = {"student_name": "Ручной Студент", "group_name": "Другая группа"}
    assert teacher.post(path + "/manual", manual).json()["status"] == "manual"
    assert teacher.post(path + "/manual", manual).status_code == 409
    assert teacher.post(path + "/close").status_code == 200
    assert attend(browser(), session).status_code == 410
    assert teacher.post(path + "/manual", manual).status_code == 410
    assert session["id"] not in teacher.get("/").text
    assert teacher.get(path + "/export").status_code == 200


def test_excel_opens_with_dates_reasons_and_safe_text(browser, teacher, session):
    assert attend(browser(), session).status_code == 200
    attend(browser(), session, name="Без Геолокации", seed="two", latitude=None)
    dangerous_name = '=HYPERLINK("https://example.invalid","click")'
    teacher.post(
        f"/api/sessions/{session['id']}/manual",
        {
            "student_name": dangerous_name,
            "group_name": "А-1",
        },
    )
    response = teacher.get(f"/api/sessions/{session['id']}/export")
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["Журнал", "Отклонённые попытки"]
    sheet = book["Журнал"]
    assert [c.value for c in sheet[1]][:5] == ["ФИО", "Группа", "Дата", "Время", "Статус"]
    assert sheet.max_row == 3
    assert sheet["A2"].value == dangerous_name and sheet["A2"].data_type == "s"
    assert sheet["C2"].is_date and sheet["D2"].is_date
    assert book.worksheets[1].max_row == 2
    assert "геолокация" in book.worksheets[1]["F2"].value


def test_qr_png_and_code_entry(teacher, browser, session):
    assert len(session["session_code"]) == 6 and session["session_code"].isdigit()
    response = teacher.get(session["qr_path"])
    image = Image.open(BytesIO(response.content))
    image.verify()
    assert image.format == "PNG"
    result = browser().post("/api/sessions/resolve", {"code": session["session_code"]})
    assert result.json()["student_url"] == session["student_url"]
    assert "ФИО" in browser().get(session["student_url"]).text
    assert session["id"] in teacher.get("/").text


def test_unknown_session_logged(client, browser):
    result = browser().post("/api/attendance/999999", student_data())
    assert result.status_code == 410
    count = client.portal.call(
        client.app.state.pool.fetchval, "SELECT count(*) FROM attendance WHERE session_id IS NULL"
    )
    assert count == 1


def test_race_accepts_only_once(browser, teacher, session):
    browsers = [browser() for _ in range(8)]
    with ThreadPoolExecutor(max_workers=8) as executor:
        results = list(executor.map(lambda b: attend(b, session).status_code, browsers))
    assert results.count(200) == 1 and results.count(409) == 7
    rows = teacher.get(f"/api/sessions/{session['id']}/attendance").json()
    assert len(rows) == 8


def test_ip_rate_limit_audited(client, browser, teacher, session):
    client.app.state.settings.ip_rate_limit = 2
    assert attend(browser(), session).status_code == 200  # Creation consumed the first request.
    result = attend(browser(), session, name="Другой Студент", seed="two")
    assert result.status_code == 429 and result.json()["reason"] == "rate_limit"
    assert len(teacher.get(f"/api/sessions/{session['id']}/attendance").json()) == 2


def test_ip_window_recovers(client, browser, session):
    client.app.state.settings.ip_rate_limit = 1
    assert attend(browser(), session).status_code == 429
    client.portal.call(
        client.app.state.pool.execute, "UPDATE ip_rate_limits SET window_start=now()-interval '2 minutes'"
    )
    assert attend(browser(), session).status_code == 200
