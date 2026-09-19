"""A real HTTP check through Nginx → Uvicorn → main PostgreSQL.

Uses its own closed demo session, never truncates the application's database.
"""

import os
from io import BytesIO
from pathlib import Path

import httpx
import pytest
from conftest import csrf_from_html, student_data
from openpyxl import load_workbook


def test_real_http_stack():
    base = os.getenv("LIVE_BASE_URL")
    if not base:
        pytest.skip("LIVE_BASE_URL not supplied")
    with (
        httpx.Client(base_url=base, timeout=15) as teacher,
        httpx.Client(base_url=base, timeout=15) as student,
    ):
        assert teacher.get("/health").json()["database"] == "ok"
        token = csrf_from_html(teacher.get("/").text)
        headers = {"X-CSRF-Token": token}
        response = teacher.post(
            "/api/sessions",
            headers=headers,
            json={
                "title": "Проверка HTTP / Nginx",
                "group_name": "ИВТ-21",
                "teacher_name": "QA",
                "campus_id": "chuvsu-g",
                "duration_minutes": 1,
                "radius": 150,
            },
        )
        assert response.status_code == 201, response.text
        session = response.json()
        try:
            assert teacher.get(session["qr_path"]).content.startswith(b"\x89PNG")
            stoken = csrf_from_html(student.get(session["student_url"]).text)
            payload = student_data()
            path = f"/api/attendance/{session['session_code']}"
            assert student.post(path, headers={"X-CSRF-Token": stoken}, json=payload).status_code == 200
            assert student.post(path, headers={"X-CSRF-Token": stoken}, json=payload).status_code == 409
            exported = teacher.get(f"/api/sessions/{session['id']}/export")
            assert load_workbook(BytesIO(exported.content))["Журнал"].max_row == 2
            if artifact_path := os.getenv("EXPORT_ARTIFACT_PATH"):
                Path(artifact_path).write_bytes(exported.content)
            assert student.get(session["teacher_url"]).status_code == 404
        finally:
            teacher.post(f"/api/sessions/{session['id']}/close", headers=headers)
