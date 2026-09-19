from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

STATUS = {"accepted": "Успешно", "manual": "Добавлен преподавателем", "rejected": "Отклонено"}


def make_excel(rows, timezone: str) -> bytes:
    """Sorted by group; separate sheets for attendance and rejected attempts."""
    book = Workbook()
    book.remove(book.active)
    for title, rejected in (("Журнал", False), ("Отклонённые попытки", True)):
        sheet = book.create_sheet(title)
        sheet.append(["ФИО", "Группа", "Дата", "Время", "Статус", "Причина"])
        for row in sorted(rows, key=lambda r: (r["group_name"], r["student_name"], r["created_at"])):
            if (row["status"] == "rejected") != rejected:
                continue
            local = row["created_at"].astimezone(ZoneInfo(timezone))
            values = [
                row["student_name"],
                row["group_name"],
                local.date(),
                local.time().replace(tzinfo=None),
                STATUS[row["status"]],
                row["reason"],
            ]
            sheet.append(values)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # User text must never become an Excel formula.
            sheet.cell(sheet.max_row, 3).number_format = "DD.MM.YYYY"
            sheet.cell(sheet.max_row, 4).number_format = "HH:MM:SS"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="163D60")
        for column, width in zip("ABCDEF", [38, 20, 16, 14, 30, 55], strict=True):
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    book.save(output)
    return output.getvalue()
